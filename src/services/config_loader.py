import hashlib
import openpyxl
from sqlalchemy.orm import Session
from src.models import Location, Category
from src.config import settings
from src.utils.logging import logger

def generate_location_id(lat: float, lng: float, pincode: str) -> str:
    s = f"{lat}{lng}{pincode}".encode('utf-8')
    return hashlib.sha256(s).hexdigest()[:12]

def generate_category_id(category_name: str) -> str:
    s = category_name.encode('utf-8')
    return hashlib.sha256(s).hexdigest()[:12]

def _geocode_pin(pincode: str, state):
    """(lat, lng) for a PIN with no coordinates in the sheet, or the reason there are none."""
    from src.services import geo_bounds, pincode_coverage

    coords = pincode_coverage._locate_or_none(pincode)
    if coords is None:
        return "No coordinates in the sheet, and the PIN code could not be found on the map"
    problem = geo_bounds.check(state, coords[0], coords[1])
    if problem:
        return f"No coordinates in the sheet, and the PIN code's lookup looks wrong: {problem}"
    return coords


def sync_config(db: Session) -> dict:
    logger.info("Starting configuration sync...")
    
    locations_added = 0
    categories_added = 0
    # Rows that could not be loaded, and why, so none disappears silently.
    skipped = []

    # 1. Load Locations
    try:
        seen_locs = set()
        wb_loc = openpyxl.load_workbook(settings.locations_file, read_only=True, data_only=True)
        sheet_loc = wb_loc["Geocoded Locations"] if "Geocoded Locations" in wb_loc.sheetnames else wb_loc.active
        
        # Skip header
        for i, row in enumerate(sheet_loc.iter_rows(values_only=True)):
            if i == 0:
                continue
            
            # The columns are: PIN Code, Latitude, Longitude, District, State, Tehsil, Anchor Village/Town
            pincode = row[0]
            lat = row[1] if len(row) > 1 else None
            lng = row[2] if len(row) > 2 else None
            district = str(row[3]).strip() if len(row) > 3 and row[3] is not None else None
            state = str(row[4]).strip() if len(row) > 4 and row[4] is not None else None
            tehsil = str(row[5]).strip() if len(row) > 5 and row[5] is not None else None
            anchor_name = str(row[6]).strip() if len(row) > 6 and row[6] is not None else None
            
            if not pincode:
                continue

            precision = None
            if lat is None or lng is None:
                # Most such rows are PINs added from the dashboard, which
                # already have coordinates in the database. Looking those up
                # again would add each a second time at a slightly different
                # point, so only a PIN the database has never seen is placed.
                known = db.query(Location).filter(Location.pincode == str(pincode).strip()).all()
                if known:
                    for existing in known:
                        if not existing.district and district: existing.district = district
                        if not existing.state and state: existing.state = state
                        if not existing.tehsil and tehsil: existing.tehsil = tehsil
                        if not existing.anchor_name and anchor_name: existing.anchor_name = anchor_name
                    continue
                # A PIN typed into the sheet without coordinates used to be
                # dropped without a word, so it never reached the server.
                # It is placed from the PIN code instead, and checked against
                # its state so a lookup that lands elsewhere is not saved.
                found = _geocode_pin(str(pincode).strip(), state)
                if isinstance(found, str):
                    skipped.append({"pincode": str(pincode).strip(), "reason": found})
                    logger.warning(f"Skipped pincode {pincode}: {found}")
                    continue
                lat, lng = found
                precision = "pin_geocoded"

            try:
                lat = float(lat)
                lng = float(lng)
            except ValueError:
                logger.warning(f"Invalid coordinates for pincode {pincode}: lat={lat}, lng={lng}")
                skipped.append({"pincode": str(pincode).strip(), "reason": "Latitude/Longitude are not numbers"})
                continue

            loc_id = generate_location_id(lat, lng, str(pincode).strip())
            
            if loc_id in seen_locs:
                continue
            seen_locs.add(loc_id)
            
            existing = db.query(Location).filter_by(location_id=loc_id).first()
            if not existing:
                loc = Location(
                    location_id=loc_id,
                    pincode=str(pincode).strip(),
                    latitude=lat,
                    longitude=lng,
                    radius_km=settings.default_radius_km,
                    district=district,
                    state=state,
                    tehsil=tehsil,
                    anchor_name=anchor_name,
                    coordinate_precision=precision,
                )
                db.add(loc)
                locations_added += 1
            else:
                # Update existing location with metadata if missing
                updated = False
                if not existing.district and district: existing.district = district; updated = True
                if not existing.state and state: existing.state = state; updated = True
                if not existing.tehsil and tehsil: existing.tehsil = tehsil; updated = True
                if not existing.anchor_name and anchor_name: existing.anchor_name = anchor_name; updated = True
                if updated:
                    db.add(existing)

    except Exception as e:
        logger.error(f"Error loading locations: {e}")
        raise

    # 2. Load Categories
    try:
        seen_cats = set()
        wb_cat = openpyxl.load_workbook(settings.categories_file, read_only=True, data_only=True)
        sheet_cat = wb_cat["Persona based Categories"] if "Persona based Categories" in wb_cat.sheetnames else wb_cat.active
        
        for i, row in enumerate(sheet_cat.iter_rows(values_only=True)):
            if i == 0:
                continue
            
            category_name = row[0]
            if not category_name:
                continue
            
            category_name = str(category_name).strip()
            cat_id = generate_category_id(category_name)
            
            if cat_id in seen_cats:
                continue
            seen_cats.add(cat_id)
            
            existing = db.query(Category).filter_by(category_id=cat_id).first()
            if not existing:
                cat = Category(
                    category_id=cat_id,
                    category_name=category_name
                )
                db.add(cat)
                categories_added += 1

    except Exception as e:
        logger.error(f"Error loading categories: {e}")
        raise

    db.commit()
    
    total_locations = db.query(Location).count()
    total_categories = db.query(Category).count()
    
    logger.info(f"Sync complete. New locations: {locations_added}, New categories: {categories_added}")

    return {
        "locations": {
            "total_in_db": total_locations,
            "new": locations_added,
            "skipped": skipped,
        },
        "categories": {
            "total_in_db": total_categories,
            "new": categories_added
        },
        "message": "Configuration synced successfully"
    }
